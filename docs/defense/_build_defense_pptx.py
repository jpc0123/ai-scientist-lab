#!/usr/bin/env python3
"""Build Scientist Lab MVP defense PPTX (product story, not a Nature paper).

Facts are frozen: baseline APS=0.0163, early_concat APS=0.0326,
same Frozen Fingerprint, ClaimGate C1 SUPPORTED. Staging 160x160 / 2ep.
"""

from __future__ import annotations

from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

OUT = Path(__file__).resolve().parent / "ScientistLab_MVP_Defense.pptx"
FONT = "Microsoft YaHei"

NAVY = RGBColor(0x0F, 0x2D, 0x3E)
INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x5A, 0x68, 0x72)
ACCENT = RGBColor(0x1F, 0x5A, 0x6E)
LINE = RGBColor(0xD6, 0xDC, 0xE0)
PALE = RGBColor(0xF4, 0xF7, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREEN = RGBColor(0x2E, 0x6B, 0x4F)
RED = RGBColor(0x8B, 0x3A, 0x3A)
GOLD = RGBColor(0x8C, 0x6C, 0x34)
PALE_GREEN = RGBColor(0xE7, 0xF0, 0xEA)
PALE_RED = RGBColor(0xF6, 0xEB, 0xEB)
PALE_GOLD = RGBColor(0xF7, 0xF1, 0xE4)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def _set_run(run, *, size: int, bold: bool = False, color=INK, name: str = FONT) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        node = rPr.find(qn(tag))
        if node is None:
            node = etree.SubElement(rPr, qn(tag))
        node.set("typeface", name)


def _fill(shape, color: RGBColor) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def _line(shape, color: RGBColor, pt: float = 1.0) -> None:
    shape.line.color.rgb = color
    shape.line.width = Pt(pt)


def add_text_box(
    slide,
    left,
    top,
    width,
    height,
    text: str,
    *,
    size: int = 16,
    bold: bool = False,
    color=INK,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP,
) -> None:
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_after = Pt(0)
    run = p.add_run()
    run.text = text
    _set_run(run, size=size, bold=bold, color=color)


def add_bullets(
    slide,
    left,
    top,
    width,
    height,
    items: list[str],
    *,
    size: int = 16,
    color=INK,
    bold_first: bool = False,
    spacing: float = 8,
) -> None:
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.level = 0
        p.space_after = Pt(spacing)
        p.space_before = Pt(0)
        run = p.add_run()
        run.text = "·  " + item
        _set_run(run, size=size, bold=(bold_first and i == 0), color=color)


def add_notes(slide, text: str) -> None:
    notes = slide.notes_slide.notes_text_frame
    notes.clear()
    p = notes.paragraphs[0]
    run = p.add_run()
    run.text = text
    _set_run(run, size=12, color=INK)


def footer(slide, page: int, total: int, source: str = "") -> None:
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(7.28), SLIDE_W, Inches(0.22))
    _fill(bar, PALE)
    label = "Scientist Lab MVP 答辩  ·  不升级声称  ·  非 SOTA"
    if source:
        label = f"{label}  ·  {source}"
    add_text_box(
        slide,
        Inches(0.45),
        Inches(7.28),
        Inches(10.6),
        Inches(0.22),
        label,
        size=10,
        color=MUTED,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    add_text_box(
        slide,
        Inches(11.4),
        Inches(7.28),
        Inches(1.5),
        Inches(0.22),
        f"{page} / {total}",
        size=10,
        color=MUTED,
        align=PP_ALIGN.RIGHT,
        anchor=MSO_ANCHOR.MIDDLE,
    )


def accent_bar(slide) -> None:
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.08), SLIDE_H)
    _fill(bar, NAVY)


def title_block(slide, title: str, kicker: str = "") -> None:
    accent_bar(slide)
    if kicker:
        add_text_box(slide, Inches(0.5), Inches(0.22), Inches(12.3), Inches(0.28), kicker, size=12, color=ACCENT, bold=True)
        add_text_box(slide, Inches(0.5), Inches(0.48), Inches(12.3), Inches(0.7), title, size=24, bold=True, color=NAVY)
    else:
        add_text_box(slide, Inches(0.5), Inches(0.32), Inches(12.3), Inches(0.7), title, size=24, bold=True, color=NAVY)
    rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(1.18), Inches(12.3), Pt(1.25))
    _fill(rule, LINE)


def card(slide, left, top, width, height, fill=PALE) -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    _fill(shape, fill)
    _line(shape, LINE, 0.75)
    shape.adjustments[0] = 0.06
    return shape


def pill(slide, left, top, width, height, text: str, fill, text_color=WHITE, size: int = 12) -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    _fill(shape, fill)
    shape.adjustments[0] = 0.5
    tf = shape.text_frame
    tf.word_wrap = True
    tf.anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    _set_run(run, size=size, bold=True, color=text_color)


def chevron_row(slide, labels: list[str], left, top, box_w, box_h, gap, fill=ACCENT) -> None:
    x = left
    for i, label in enumerate(labels):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, top, box_w, box_h)
        _fill(shape, fill)
        shape.adjustments[0] = 0.12
        tf = shape.text_frame
        tf.word_wrap = True
        tf.anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = label
        _set_run(run, size=12, bold=True, color=WHITE)
        if i < len(labels) - 1:
            arrow = slide.shapes.add_shape(
                MSO_SHAPE.RIGHT_ARROW,
                x + box_w + Inches(0.04),
                top + (box_h - Inches(0.16)) / 2,
                gap - Inches(0.08),
                Inches(0.16),
            )
            _fill(arrow, LINE)
        x += box_w + gap


def build() -> Path:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]
    total = 12

    # --- 1 Cover ---
    s = prs.slides.add_slide(blank)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), SLIDE_W, SLIDE_H)
    _fill(bg, NAVY)
    stripe = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.18), SLIDE_H)
    _fill(stripe, ACCENT)
    add_text_box(s, Inches(0.7), Inches(1.35), Inches(11.8), Inches(0.35), "Scientist Lab  ·  MVP 答辩", size=14, color=RGBColor(0xA8, 0xC4, 0xD0), bold=True)
    add_text_box(s, Inches(0.7), Inches(1.85), Inches(11.8), Inches(1.3), "这是一套规则/状态驱动的\n可控自主实验系统", size=34, bold=True, color=WHITE)
    add_text_box(
        s,
        Inches(0.7),
        Inches(3.4),
        Inches(11.8),
        Inches(0.9),
        "当前交付不是完整 LLM AI Scientist，也不是更好的检测器。\n评委要看到的是：系统能自主迭代，并且能限制自己根据证据说到什么程度。",
        size=16,
        color=RGBColor(0xD7, 0xE4, 0xEA),
    )
    pill(s, Inches(0.7), Inches(4.55), Inches(3.35), Inches(0.38), "不是 APS 涨点叙事", NAVY, RGBColor(0xA8, 0xC4, 0xD0), 12)
    pill(s, Inches(4.2), Inches(4.55), Inches(2.55), Inches(0.38), "不是 SOTA", NAVY, RGBColor(0xA8, 0xC4, 0xD0), 12)
    pill(s, Inches(6.9), Inches(4.55), Inches(3.7), Inches(0.38), "early_concat ≠ FDPN", NAVY, RGBColor(0xA8, 0xC4, 0xD0), 12)
    add_text_box(
        s,
        Inches(0.7),
        Inches(5.35),
        Inches(11.8),
        Inches(0.9),
        "Freeze tag  mvp-freeze-m1-m4-claimgate-c1\nSHA  130b02cfbb5521829e959d10b99d17fd5fff28ab    ·    2026-08-17",
        size=13,
        color=RGBColor(0x8F, 0xB0, 0xBC),
    )
    add_notes(
        s,
        "开场先定调。不要把系统讲成「又一个会训检测器的 Agent」，也不要讲成「已经接了 GPT 的完整 AI Scientist」。"
        "两句拆开：现在能自主跑实验闭环；科研认知仍由人+GPT 在系统外提供。下一代才把 Planner/Reviewer 经 LLM API 搬进系统。"
        "全程不要追求 APS 涨点叙事。",
    )

    # --- 2 是什么 / 不是什么 ---
    s = prs.slides.add_slide(blank)
    title_block(s, "当前 MVP 管过程与声称，不假装自己是完整 AI Scientist", "产品定位")
    card(s, Inches(0.5), Inches(1.4), Inches(6.0), Inches(5.55), PALE_GREEN)
    card(s, Inches(6.8), Inches(1.4), Inches(6.0), Inches(5.55), PALE_RED)
    add_text_box(s, Inches(0.75), Inches(1.55), Inches(5.5), Inches(0.4), "是什么", size=18, bold=True, color=GREEN)
    add_text_box(s, Inches(7.05), Inches(1.55), Inches(5.5), Inches(0.4), "不是什么", size=18, bold=True, color=RED)
    add_bullets(
        s,
        Inches(0.75),
        Inches(2.05),
        Inches(5.5),
        Inches(4.6),
        [
            "人设目标、边界、预算；系统按状态机跑完一轮并写出下一轮计划",
            "行为闭环：Gate → Run → Evidence → Rubric → Memory → Next Plan",
            "声称闭环：ClaimGate 限制证据最多支持说到哪一步",
            "研发阶段科研认知由「人 + GPT」协作提供；系统负责执行与守门",
        ],
        size=15,
        color=INK,
        spacing=14,
    )
    add_bullets(
        s,
        Inches(7.05),
        Inches(2.05),
        Inches(5.5),
        Inches(4.6),
        [
            "完整 LLM AI Scientist（Planner / Reviewer 尚未经 LLM API 进系统）",
            "更好的检测器，或论文级 SOTA 结果",
            "FDPN / C2；early_concat 只是已有 staging 混合",
            "640×640 / 20 epoch 论文协议（本 C1 是 160×160 / 2 epoch staging）",
        ],
        size=15,
        color=INK,
        spacing=14,
    )
    footer(s, 2, total, "整理自 MVP_RELEASE_NOTES / 设计架构")
    add_notes(
        s,
        "左边四条是产品承诺，右边四条是现场禁语。若评委问「你们是不是做了一个更强的 RGB-T 检测器」，直接回答：不是。"
        "检测任务只是闭环的执行域。KEEP 也不等于模块有效。",
    )

    # --- 3 评委要看到什么 ---
    s = prs.slides.add_slide(blank)
    title_block(s, "评委要看到的不是涨点，而是「能迭代」且「能闭嘴」", "答辩核心")
    card(s, Inches(0.5), Inches(1.45), Inches(6.0), Inches(3.35), PALE)
    card(s, Inches(6.8), Inches(1.45), Inches(6.0), Inches(3.35), PALE)
    add_text_box(s, Inches(0.75), Inches(1.6), Inches(5.5), Inches(0.4), "1  系统能自主迭代", size=18, bold=True, color=ACCENT)
    add_text_box(
        s,
        Inches(0.75),
        Inches(2.15),
        Inches(5.5),
        Inches(2.3),
        "在固定协议与预算内，Manager 按状态机调度：计划获批、真实执行、读数、判断、写入记忆、提出下一轮。负结果可被引用并改向，不是 error。",
        size=15,
        color=INK,
    )
    add_text_box(s, Inches(7.05), Inches(1.6), Inches(5.5), Inches(0.4), "2  系统能限制自己说到哪一步", size=18, bold=True, color=ACCENT)
    add_text_box(
        s,
        Inches(7.05),
        Inches(2.15),
        Inches(5.5),
        Inches(2.3),
        "ClaimGate 是确定性规则服务，不是 Agent。KEEP 不能升格为 SUPPORTED。probe、mAP、单臂对照都会被挡住。证据不够时，系统必须闭嘴。",
        size=15,
        color=INK,
    )
    warn = card(s, Inches(0.5), Inches(5.0), Inches(12.3), Inches(1.95), PALE_GOLD)
    add_text_box(s, Inches(0.75), Inches(5.15), Inches(11.8), Inches(0.35), "本场不要讲成", size=14, bold=True, color=GOLD)
    add_text_box(
        s,
        Inches(0.75),
        Inches(5.55),
        Inches(11.8),
        Inches(1.2),
        "APS 从 0.0163 涨到 0.0326 所以产品成功；SOTA；把 676 tests 堆成正文。测试是核对手段，故事是两条闭环与声称纪律。",
        size=15,
        color=INK,
    )
    footer(s, 3, total)
    add_notes(
        s,
        "这页是整场答辩的脊柱。后面所有页都服务这两句：能迭代、能限制声称。"
        "676 passed 只在被问到验证口径时口头提一句，不要展开测试清单。",
    )

    # --- 4 两条闭环 ---
    s = prs.slides.add_slide(blank)
    title_block(s, "冻结的是两条闭环，不是一条训练流水线", "架构 → 代码")
    add_text_box(s, Inches(0.5), Inches(1.4), Inches(12.3), Inches(0.35), "自主科研行为闭环", size=16, bold=True, color=NAVY)
    chevron_row(
        s,
        ["Gate", "Run", "Evidence", "Rubric", "Memory", "Next Plan"],
        Inches(0.5),
        Inches(1.85),
        Inches(1.7),
        Inches(0.62),
        Inches(0.32),
        ACCENT,
    )
    add_text_box(
        s,
        Inches(0.5),
        Inches(2.6),
        Inches(12.3),
        Inches(0.4),
        "源码：core/manager.py  +  GateEngine / ResultParser / EvidenceValidator / DecisionRubric / MemoryWriter / next_plan",
        size=12,
        color=MUTED,
    )
    add_text_box(s, Inches(0.5), Inches(3.2), Inches(12.3), Inches(0.35), "科学声称闭环", size=16, bold=True, color=NAVY)
    chevron_row(
        s,
        ["Evidence", "Claim 草案", "ClaimGate", "SUPPORTED", "INCONCLUSIVE", "BLOCKED"],
        Inches(0.5),
        Inches(3.65),
        Inches(1.7),
        Inches(0.62),
        Inches(0.32),
        GREEN,
    )
    add_text_box(
        s,
        Inches(0.5),
        Inches(4.4),
        Inches(12.3),
        Inches(0.4),
        "源码：core/claim_gate.py  ·  确定性规则服务，不是 Agent，不覆盖 Reviewer 的 KEEP / DISCARD",
        size=12,
        color=MUTED,
    )
    card(s, Inches(0.5), Inches(5.0), Inches(12.3), Inches(1.95), PALE)
    add_bullets(
        s,
        Inches(0.7),
        Inches(5.15),
        Inches(11.9),
        Inches(1.65),
        [
            "Architecture-to-Code Audit：PASS（checklist 9/9）。权威文本：docs/ARCHITECTURE_TO_CODE_AUDIT.md",
            "KEEP / DISCARD ≠ SUPPORTED / INCONCLUSIVE / BLOCKED；scientific_outcome 是只读投影，不是第二条声称状态机",
            "人管制度，Manager 管过程；确定性问题交给系统，模糊科研判断才交给未来的 LLM",
        ],
        size=14,
        spacing=8,
    )
    footer(s, 4, total, "整理自 设计架构.md / MVP_FREEZE.md")
    add_notes(
        s,
        "画两条闭环时用手指出：上面一条证明系统会干活并学习下一轮；下面一条证明系统不会乱说话。"
        "ClaimGate 与 Reviewer 并存：Reviewer 决定 KEEP/DISCARD，ClaimGate 决定声称强度。二者不是同一个开关。",
    )

    # --- 5 M1-M4 ---
    s = prs.slides.add_slide(blank)
    title_block(s, "M1–M4 证明的是可审计的自主实验行为，不是模型变强", "行为证据")
    items = [
        ("Gate 先于点火", "未批准的计划不能执行。fingerprint 变了会 BLOCK。REAL 必须 --execute；doctor 未就绪不得伪造 metrics。"),
        ("VALID 才进入判断", "INVALID / FAILED 是工程或协议错误。科研负结果仍可 DISCARD，并可被下一轮引用。"),
        ("Memory → Next Plan", "Round≥1 必须引用已写入的 lesson / strategy。空引用或假引用会被 Gate 拒绝。"),
        ("Git 与决策绑定", "KEEP 才更新 best_sha；DISCARD 切回 best，禁止 reset --hard。无 VALID 不 apply。"),
    ]
    for i, (head, body) in enumerate(items):
        col = i % 2
        row = i // 2
        left = Inches(0.5) + col * Inches(6.35)
        top = Inches(1.42) + row * Inches(2.55)
        card(s, left, top, Inches(6.1), Inches(2.35), PALE)
        pill(s, left + Inches(0.2), top + Inches(0.22), Inches(0.42), Inches(0.32), str(i + 1), ACCENT, WHITE, 12)
        add_text_box(s, left + Inches(0.75), top + Inches(0.2), Inches(5.1), Inches(0.4), head, size=16, bold=True, color=NAVY)
        add_text_box(s, left + Inches(0.25), top + Inches(0.75), Inches(5.6), Inches(1.35), body, size=14, color=INK)
    footer(s, 5, total, "行为证据在测试与 .run 产物；676 tests 不是正文")
    add_notes(
        s,
        "不要把这一页讲成里程碑流水账。四块各举一个现场可指的事实："
        "REAL 必须 --execute；probe 的 APS=0 不是 SOTA 也不是 C1；N+1 必须 cite memory；"
        "Manager 单臂 ClaimGate BLOCKED 是正确行为。被问到测试时，指向 test_claim_gate_m4.py 与 test_manager_multround_m4.py。",
    )

    # --- 6 ClaimGate ---
    s = prs.slides.add_slide(blank)
    title_block(s, "ClaimGate 让系统先判断「能不能说」，而不是先写结论", "声称纪律")
    rows = [
        ("KEEP ≠ SUPPORTED", "Reviewer 决定保留工作树，不决定科学声称。"),
        ("DISCARD ≠ 模块无效", "负结果可引用、可改向，不能写成「融合没用」。"),
        ("probe ≠ formal", "16/8 fast_eval 禁止当 C1。probe 协议 allow_scientific_claims=false。"),
        ("mAP ≠ APS", "声称 small-object 但证据只有 mAP → BLOCKED。不得从 mAP 发明 APS。"),
        ("单臂 → BLOCKED", "只有 baseline_metrics、无配对 Frozen Fingerprint 时，C1 必须挡住。这是正确行为。"),
        ("配对才可能 C1", "CLI：claim-gate --baseline-run-dir。同一 fingerprint 上的 formal 对照。"),
    ]
    for i, (head, body) in enumerate(rows):
        col = i % 2
        row = i // 2
        left = Inches(0.5) + col * Inches(6.35)
        top = Inches(1.4) + row * Inches(1.75)
        fill = PALE_RED if "≠" in head or "BLOCKED" in head else PALE
        card(s, left, top, Inches(6.1), Inches(1.58), fill)
        add_text_box(s, left + Inches(0.25), top + Inches(0.18), Inches(5.6), Inches(0.4), head, size=16, bold=True, color=NAVY)
        add_text_box(s, left + Inches(0.25), top + Inches(0.62), Inches(5.6), Inches(0.75), body, size=14, color=INK)
    footer(s, 6, total, "core/claim_gate.py")
    add_notes(
        s,
        "现场把「KEEP 了所以融合有效」当成最危险的误读。指出 claim_gate_c1.json 里 keep_is_not_claim=true，"
        "reason 写明 KEEP/DISCARD 未决定声称。再打开 candidate_resume/claim_gate.json 的 BLOCKED，强调这不是 bug。",
    )

    # --- 7 Formal C1 数字 ---
    s = prs.slides.add_slide(blank)
    title_block(s, "Formal C1 只证明：同一 fingerprint 下，staging 对照可被 C1 声称", "已接受事实  ·  数字不得改")
    table_shape = s.shapes.add_table(4, 4, Inches(0.5), Inches(1.45), Inches(12.3), Inches(2.55))
    table = table_shape.table
    table.columns[0].width = Inches(2.4)
    table.columns[1].width = Inches(3.5)
    table.columns[2].width = Inches(3.3)
    table.columns[3].width = Inches(3.1)
    headers = ["臂", "HOW", "APS（四位口径）", "角色"]
    data = [
        ["Formal-01 baseline", "rgb + none", "0.0163", "matched baseline"],
        ["Formal-02 candidate", "rgbt + early_concat", "0.0326", "C1 对照臂"],
        ["配对 ClaimGate", "同一 Frozen Fingerprint", "C1 SUPPORTED", "--baseline-run-dir"],
    ]
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.text = ""
        p = cell.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        run = p.add_run()
        run.text = h
        _set_run(run, size=13, bold=True, color=WHITE)
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
    for i, row in enumerate(data, start=1):
        for j, val in enumerate(row):
            cell = table.cell(i, j)
            cell.text = ""
            p = cell.text_frame.paragraphs[0]
            run = p.add_run()
            run.text = val
            bold = j == 2
            color = GREEN if (i == 3 and j == 2) else INK
            _set_run(run, size=14, bold=bold, color=color)
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if i % 2 else PALE
    add_text_box(
        s,
        Inches(0.5),
        Inches(4.15),
        Inches(12.3),
        Inches(0.7),
        "handle 内 raw：APS=0.016313298031160568 / 0.03256971555632778（AP_small 同值）。fingerprint_id=FP-RGBT-DFINE-FORMAL-C1-V1；HOW 只进 notes，不进哈希。",
        size=13,
        color=MUTED,
    )
    card(s, Inches(0.5), Inches(4.9), Inches(12.3), Inches(2.05), PALE_GOLD)
    add_text_box(s, Inches(0.75), Inches(5.05), Inches(11.8), Inches(0.35), "必须同时说出口的限制", size=14, bold=True, color=GOLD)
    add_text_box(
        s,
        Inches(0.75),
        Inches(5.45),
        Inches(11.8),
        Inches(1.3),
        "160×160 / 2 epoch 是 staging / Formal-C1，不是 640/20ep 论文协议。early_concat 是仓库已有 staging 混合，不是 FDPN，不得升 C2。"
        "Manager 单臂 claim_gate.json = BLOCKED，正确。scientific_outcome=INCONCLUSIVE 与 ClaimGate SUPPORTED 可并存。",
        size=14,
        color=INK,
    )
    footer(s, 7, total, "整理自 MVP_FREEZE.md；.run/ 不进 git")
    add_notes(
        s,
        "读表时先说「同一 fingerprint」，再说两个 APS，最后说配对 SUPPORTED。"
        "不要停在「涨了一倍」这种话。立刻补限制：staging、不是 FDPN、单臂 BLOCKED 是正确的。"
        "产物在 .run/formal_c1_aps_early_concat/，gitignore，本机才有。",
    )

    # --- 8 能说 / 不能说 ---
    s = prs.slides.add_slide(blank)
    title_block(s, "这组数字允许说 C1，不允许说更好、SOTA 或 FDPN", "声称边界")
    card(s, Inches(0.5), Inches(1.4), Inches(6.0), Inches(5.55), PALE_GREEN)
    card(s, Inches(6.8), Inches(1.4), Inches(6.0), Inches(5.55), PALE_RED)
    add_text_box(s, Inches(0.75), Inches(1.55), Inches(5.5), Inches(0.4), "能说", size=18, bold=True, color=GREEN)
    add_text_box(s, Inches(7.05), Inches(1.55), Inches(5.5), Inches(0.4), "不能说", size=18, bold=True, color=RED)
    add_bullets(
        s,
        Inches(0.75),
        Inches(2.1),
        Inches(5.5),
        Inches(4.5),
        [
            "在同一 Frozen Fingerprint 上，formal candidate 的 APS 高于 matched baseline",
            "配对 ClaimGate 将 C1 判为 SUPPORTED，且 KEEP 未决定该声称",
            "系统能自主完成 Gate→Run→Evidence→Memory→Next Plan",
            "系统能在证据不足时 BLOCKED / INCONCLUSIVE，而不是编造指标",
        ],
        size=15,
        spacing=12,
    )
    add_bullets(
        s,
        Inches(7.05),
        Inches(2.1),
        Inches(5.5),
        Inches(4.5),
        [
            "产品是更好的检测器，或我们刷到了 SOTA",
            "early_concat = FDPN，或已经证明融合模块有效（那是 C2）",
            "这是 640/20ep 论文协议结果；probe 的 16/8 也算 C1",
            "单臂 BLOCKED 是 C1 失败；INCONCLUSIVE 投影否定了 ClaimGate",
        ],
        size=15,
        spacing=12,
    )
    footer(s, 8, total)
    add_notes(
        s,
        "这页用来挡追问。若评委说「那你们的贡献是涨点」，拉回左边第三条和第四条：贡献是可控自主实验与声称守门。"
        "C2 需要 ablation，本 MVP 没有，所以不能说模块有效。",
    )

    # --- 9 Demo ---
    s = prs.slides.add_slide(blank)
    title_block(s, "现场 Demo 走真实命令，不现场发明指标，不重跑 Formal GPU", "Demo")
    steps = [
        ("1  核对身份", "git rev-parse mvp-freeze-m1-m4-claimgate-c1^{}\n期望 SHA 130b02cf…"),
        ("2  dry-run 状态机", "manager-run 默认不点火。看 Gate→编排，不写伪造 metrics。"),
        ("3  打开已有 C1 产物", ".run/formal_c1_aps_early_concat/（gitignore，本机才有）"),
        ("4  重放 ClaimGate", "claim-gate --baseline-run-dir … 无 GPU。对比单臂 BLOCKED。"),
    ]
    for i, (head, body) in enumerate(steps):
        left = Inches(0.5) + i * Inches(3.18)
        card(s, left, Inches(1.4), Inches(3.05), Inches(2.7), PALE)
        add_text_box(s, left + Inches(0.15), Inches(1.55), Inches(2.75), Inches(0.7), head, size=15, bold=True, color=NAVY)
        add_text_box(s, left + Inches(0.15), Inches(2.25), Inches(2.75), Inches(1.6), body, size=12, color=INK)
    card(s, Inches(0.5), Inches(4.3), Inches(12.3), Inches(2.65), PALE_GOLD)
    add_text_box(s, Inches(0.75), Inches(4.45), Inches(11.8), Inches(0.35), "命令红线", size=14, bold=True, color=GOLD)
    add_text_box(
        s,
        Inches(0.75),
        Inches(4.9),
        Inches(11.8),
        Inches(1.8),
        "REAL 必须加 --execute；可加 --require-live-ready，doctor 为 false 时拒绝 GPU 且不得伪造 metrics。\n"
        "probe ≠ formal：probe 协议禁止科学声称；Formal C1 已跑完，答辩不要现场重跑 GPU。\n"
        "完整逐步命令见 docs/defense/DEMO_SCRIPT.md。",
        size=14,
        color=INK,
    )
    footer(s, 9, total, "DEMO_SCRIPT.md")
    add_notes(
        s,
        "Demo 建议 8–10 分钟：身份核对 1 分钟，dry-run 2 分钟，打开 metrics/handle/claim_gate 4 分钟，对照单臂 BLOCKED 2 分钟。"
        "若 .run 目录不在这台机器上，不要临时造数字，改为只演示 dry-run + git 内协议/schema。",
    )

    # --- 10 LLM ---
    s = prs.slides.add_slide(blank)
    title_block(s, "下一代才把 Planner / Reviewer 经 LLM API 搬进系统", "人 + GPT 现在在哪")
    phases = [
        ("现在（本 MVP）", ACCENT, "Planner / Reviewer 是规则优先实现。科研假设、协议与声称草案由人 + GPT 在系统外协作产生。系统负责 Gate、执行、证据、记忆与声称守门。"),
        ("下一代", GREEN, "把 Planner / Reviewer 接到 LLM API：开放式科研判断进系统，但仍走同一套 Protocol、Fingerprint、Gate、ClaimGate。LLM 不覆盖确定性守门。"),
        ("不这样做", RED, "不把 LLM 当成可以改 APS、跳过 Gate、或把 KEEP 写成 SUPPORTED 的法官。不把 v2.5 / Web / dataset 战役讲成已交付 MVP。"),
    ]
    for i, (head, color, body) in enumerate(phases):
        top = Inches(1.4) + i * Inches(1.75)
        card(s, Inches(0.5), top, Inches(12.3), Inches(1.6), PALE)
        pill(s, Inches(0.75), top + Inches(0.22), Inches(2.5), Inches(0.36), head, color, WHITE, 12)
        add_text_box(s, Inches(3.45), top + Inches(0.2), Inches(9.05), Inches(1.2), body, size=15, color=INK)
    footer(s, 10, total, "整理自 设计架构.md：确定性问题交给系统")
    add_notes(
        s,
        "被问「为什么现在没有 GPT 在环内」时：这是有意的阶段划分，不是能力没写完就假装写完。"
        "架构原则是确定性问题交给系统，模糊科研判断才交给 LLM。ClaimGate、Fingerprint、Gate 必须先硬，再接 LLM。",
    )

    # --- 11 边界 ---
    s = prs.slides.add_slide(blank)
    title_block(s, "未进入 MVP 的，现场不要讲成已交付", "能力边界")
    excluded = [
        "新 GPU probe / formal 重跑",
        "trajectory_step / 训练飞轮",
        "preference / SFT / DPO / RL",
        "Bounded Tree",
        "LLM Planner / Reviewer（API）",
        "FDPN C2",
        "论文级 SOTA",
        "Web 产品化",
        "v2.5 GATE 战役",
        "dataset 注册战役",
    ]
    for i, label in enumerate(excluded):
        col = i % 5
        row = i // 5
        left = Inches(0.5) + col * Inches(2.52)
        top = Inches(1.45) + row * Inches(1.35)
        card(s, left, top, Inches(2.38), Inches(1.18), PALE_RED)
        add_text_box(s, left + Inches(0.12), top + Inches(0.28), Inches(2.14), Inches(0.65), label, size=13, bold=True, color=RED, align=PP_ALIGN.CENTER)
    card(s, Inches(0.5), Inches(4.35), Inches(12.3), Inches(2.6), PALE)
    add_bullets(
        s,
        Inches(0.7),
        Inches(4.5),
        Inches(11.9),
        Inches(2.25),
        [
            "当前工作树里的 v2.5 脏文件是 post-MVP worktree state，开发线 feat/post-mvp-v25，不得回灌 freeze tag",
            "不要 git tag -f，不要 amend freeze commit，不要把 .run / datasets / web / llm 打进该 tag",
            "保护分支 freeze/mvp-m1-m4-claimgate-c1 与 tag 同 SHA；后续开发从冻结点之后继续",
        ],
        size=15,
        spacing=10,
    )
    footer(s, 11, total, "MVP_FREEZE.md / MVP_RELEASE_NOTES.md")
    add_notes(
        s,
        "如果屏幕上出现 v2.5 目录或 Web UI，主动说明：那是封版后的工作树，不是今天答辩的交付物。"
        "今天只认 freeze tag 指向的 M1–M4 + ClaimGate + Formal C1。",
    )

    # --- 12 收束 ---
    s = prs.slides.add_slide(blank)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), SLIDE_W, SLIDE_H)
    _fill(bg, NAVY)
    stripe = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.18), SLIDE_H)
    _fill(stripe, ACCENT)
    add_text_box(s, Inches(0.7), Inches(1.2), Inches(11.8), Inches(0.35), "带走三句话", size=14, color=RGBColor(0xA8, 0xC4, 0xD0), bold=True)
    add_text_box(
        s,
        Inches(0.7),
        Inches(1.65),
        Inches(11.8),
        Inches(1.5),
        "交付的是可控自主实验系统。\n它能迭代，也能限制自己说到哪一步。",
        size=28,
        bold=True,
        color=WHITE,
    )
    lines = [
        "行为闭环已冻结：Gate → Run → Evidence → Rubric → Memory → Next Plan",
        "声称闭环已冻结：ClaimGate；KEEP ≠ Claim；C1 仅在配对 fingerprint 上 SUPPORTED",
        "口径：baseline APS=0.0163，early_concat APS=0.0326；160×160/2ep staging，不是论文协议，不是 FDPN",
    ]
    add_bullets(s, Inches(0.7), Inches(3.4), Inches(11.8), Inches(2.0), lines, size=16, color=RGBColor(0xD7, 0xE4, 0xEA), spacing=10)
    add_text_box(
        s,
        Inches(0.7),
        Inches(5.5),
        Inches(11.8),
        Inches(1.1),
        "tag  mvp-freeze-m1-m4-claimgate-c1\nSHA  130b02cfbb5521829e959d10b99d17fd5fff28ab",
        size=14,
        color=RGBColor(0x8F, 0xB0, 0xBC),
    )
    add_notes(
        s,
        "收束不要加新数字。若还有时间，回到 Demo 文档和证据链一页纸。"
        "问「下一步」时只说：从 freeze 之后接 LLM Planner/Reviewer，不推翻 Gate 与 ClaimGate。不要现场承诺 v2.5 时间表。",
    )

    prs.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print(path)
    print("slides", 12)
